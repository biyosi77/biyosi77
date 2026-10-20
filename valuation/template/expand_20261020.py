"""Reproduce the 133-stock template from the untouched 90-stock YoY workbook."""
from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from valuation.workbook import yoy_formula


HERE = Path(__file__).parent
BASE = HERE / "2027_valuation_20261016_yoy.xlsx"
ADDITIONS = HERE / "additions_20261020.json"
OUTPUT = HERE / "2027_valuation_20261020_133.xlsx"


def extend_sheet(sheet, source_row: int, target_row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        source, target = sheet.cell(source_row, col), sheet.cell(target_row, col)
        target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def main() -> None:
    additions = json.loads(ADDITIONS.read_text(encoding="utf-8"))["stocks"]
    if len(additions) != 43 or len({item["code"] for item in additions}) != 43:
        raise ValueError("expected 43 unique new stocks")
    workbook = load_workbook(BASE)
    main_sheet, raw = workbook["2027 밸류에이션"], workbook["원자료"]
    original = {str(raw.cell(row, 2).value) for row in range(10, 100)}
    if len(original) != 90 or any(item["code"] in original for item in additions):
        raise ValueError("the 90-stock base or addition mapping changed")
    for row, item in enumerate(additions, start=100):
        extend_sheet(main_sheet, 99, row, 17)
        extend_sheet(raw, 99, row, 24)
        code, name = item["code"], item["name"]
        raw.cell(row, 1).value = name
        raw.cell(row, 2).value = code
        main_sheet.cell(row, 1).value = "미산출"
        main_sheet.cell(row, 2).value = name
        main_sheet.cell(row, 3).value = code
        main_sheet.cell(row, 4).value = f'=IF(ISNUMBER(원자료!I{row}),원자료!I{row},"미확인")'
        main_sheet.cell(row, 5).value = f'=IF(AND(ISNUMBER(D{row}),ISNUMBER(원자료!J{row})),D{row}*원자료!J{row}/100000000,"미산출")'
        for col, raw_col in ((6, "T"), (7, "D"), (9, "U"), (10, "E"), (14, "O")):
            main_sheet.cell(row, col).value = f'=IF(ISNUMBER(원자료!{raw_col}{row}),원자료!{raw_col}{row},"미확인")'
        # The recurring report builder replaces these formulas and fills cached results.
        main_sheet.cell(row, 8).value = yoy_formula("F", "G", row)
        main_sheet.cell(row, 11).value = yoy_formula("I", "J", row)
        main_sheet.cell(row, 12).value = f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(G{row}),G{row}>0),E{row}/G{row},"미산출")'
        main_sheet.cell(row, 13).value = f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(J{row}),J{row}>0),E{row}/J{row},"미산출")'
        main_sheet.cell(row, 17).value = "추가 / 추정치 미제공"
    last = 99 + len(additions)
    for sheet, table_name, last_col in ((main_sheet, "ValuationTable", "Q"),
                                        (raw, "SourceTable", "X")):
        table = sheet.tables[table_name]
        table.ref = f"A9:{last_col}{last}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    workbook.save(OUTPUT)
    print(f"{OUTPUT}: {last - 9} stocks")


if __name__ == "__main__":
    main()
