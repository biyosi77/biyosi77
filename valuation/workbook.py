"""Preserve the supplied workbook while refreshing code-joined valuation rows."""
from __future__ import annotations

from copy import copy
from math import isclose
from datetime import date
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.table import TableColumn
from openpyxl.utils import get_column_letter

from .api import Quote

FIRST, LAST = 10, 99
MAIN, RAW = "2027 밸류에이션", "원자료"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def records(template):
    workbook = load_workbook(template)
    if workbook.sheetnames != [MAIN, RAW]:
        raise ValueError("template sheet structure changed")
    source, main = workbook[RAW], workbook[MAIN]
    codes = {}
    for row in range(FIRST, LAST + 1):
        code = str(source.cell(row, 2).value or "")
        if len(code) != 6 or not code.isdigit() or code in codes:
            raise ValueError("source stock code missing or duplicated")
        codes[code] = row
    if len(codes) != 90 or {str(main.cell(r, 3).value) for r in range(FIRST, LAST + 1)} != set(codes):
        raise ValueError("template must contain the same 90 unique codes on both sheets")
    return workbook, codes


def metric(cap, profit):
    return cap / profit if isinstance(cap, (int, float)) and isinstance(profit, (int, float)) and profit > 0 else None


def yoy(previous, current):
    """Match the supplied workbook's profit-growth treatment, including turnarounds."""
    if not isinstance(previous, (int, float)) or not isinstance(current, (int, float)):
        return "미확인"
    if previous == 0:
        return "흑자전환" if current > 0 else "손익분기" if current == 0 else "미확인"
    if previous < 0:
        return "흑자전환" if current > 0 else "손익분기" if current == 0 else "적자지속"
    return "적자전환" if current < 0 else -1 if current == 0 else current / previous - 1


def yoy_formula(previous, current, row):
    p, c = f"{previous}{row}", f"{current}{row}"
    return (f'=IF(AND(ISNUMBER({p}),ISNUMBER({c})),IF({p}=0,'
            f'IF({c}>0,"흑자전환",IF({c}=0,"손익분기","미확인")),'
            f'IF({p}<0,IF({c}>0,"흑자전환",IF({c}=0,"손익분기","적자지속")),'
            f'IF({c}<0,"적자전환",IF({c}=0,-1,{c}/{p}-1)))),"미확인")')


def _snapshot(cell):
    return (cell.value, copy(cell._style), copy(cell.comment), copy(cell.hyperlink))


def _restore(cell, snapshot):
    cell.value, cell._style, cell.comment, cell._hyperlink = snapshot


def _set_cached_values(path, cache):
    """Populate OOXML formula caches so viewers also see verified results without Excel."""
    path = Path(path)
    with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as tmp:
        temp = Path(tmp.name)
    try:
        with ZipFile(path) as source, ZipFile(temp, "w") as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(content)
                    for cell in root.findall(".//m:sheetData/m:row/m:c", NS):
                        address = cell.get("r")
                        if address not in cache or cell.find("m:f", NS) is None:
                            continue
                        value = cache[address]
                        node = cell.find("m:v", NS)
                        if node is None:
                            node = ET.SubElement(cell, "{%s}v" % NS["m"])
                        if isinstance(value, str):
                            cell.set("t", "str")
                        else:
                            cell.attrib.pop("t", None)
                        node.text = str(value)
                    content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(info, content)
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def build(template, output, day: date, slot: str, quotes: dict[str, Quote], *, test=False):
    workbook, rows = records(template)
    main, source = workbook[MAIN], workbook[RAW]
    if set(quotes) != set(rows):
        raise ValueError("quote mapping must include each of the 90 codes exactly once")
    old_main = {str(main.cell(r, 3).value): [_snapshot(main.cell(r, c)) for c in range(1, 18)]
                for r in range(FIRST, LAST + 1)}
    profits = {code: (source.cell(r, 20).value, source.cell(r, 4).value,
                      source.cell(r, 21).value, source.cell(r, 5).value,
                      source.cell(r, 15).value) for code, r in rows.items()}
    cap = {code: q.price * q.shares / 100_000_000 if q.price is not None and q.shares is not None else None
           for code, q in quotes.items()}
    per = {code: metric(cap[code], profits[code][1]) for code in rows}
    por = {code: metric(cap[code], profits[code][3]) for code in rows}
    ordered = sorted(rows, key=lambda c: (per[c] is None,
                                         per[c] if per[c] is not None else float("inf"), c))
    mmdd = day.strftime("%m/%d")
    label = mmdd + (" 11:30 주가" if slot == "1130" else " 종가")
    main["D9"], source["I9"] = label, label
    main["A3"] = f"{day.isoformat()} {'11:30' if slot == '1130' else 'KRX 정규장 종가'} 기준(KST). 금액: 억원, 주가: 원, 배수: 배."
    main["A5"] = "PER = 시가총액 / 27E 순이익. POR = 시가총액 / 27E 영업이익. 시가총액 = 당일 기준 주가 × 확인된 상장주식 수. 26E·27E 추정치와 YoY 기준은 원본을 따름."
    source["A4"] = f"{day.isoformat()} {label} 기준. KRX 정규장만 사용하며 NXT·시간외·전일 가격은 제외."
    if test:
        main["A3"] = "테스트 / 실제 시세 검증 없음 — " + main["A3"].value
    source["Y9"], source["Z9"], source["AA9"], source["AB9"] = (
        "가격 출처", "가격 기준시각", "조회시각", "예외 사유")
    for code, row in rows.items():
        q = quotes[code]
        source.cell(row, 9).value = q.price if q.price is not None else "미확인"
        source.cell(row, 10).value = q.shares if q.shares is not None else "미확인"
        source.cell(row, 11).value = q.basis or "미확인"
        source.cell(row, 13).value = q.source
        source.cell(row, 14).value = q.source if slot == "close" else "해당 없음"
        for col, value in ((25, q.source), (26, q.basis or "미확인"),
                           (27, q.checked_at), (28, q.reason or "")):
            source.cell(row, col).value = value
            source.cell(row, col)._style = copy(source.cell(row, 13)._style)
    source.column_dimensions["Y"].width = 36
    source.column_dimensions["Z"].width = 26
    source.column_dimensions["AA"].width = 28
    source.column_dimensions["AB"].width = 38
    table = source.tables["SourceTable"]
    table.ref = "A9:AB99"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref
    for column in range(25, 29):
        table.tableColumns.append(TableColumn(id=column,
            name=str(source.cell(9, column).value)))
    cache = {}
    ranked = sum(per[c] is not None for c in ordered)
    for idx, code in enumerate(ordered):
        row, raw_row = FIRST + idx, rows[code]
        for col, snapshot in enumerate(old_main[code], 1):
            _restore(main.cell(row, col), snapshot)
        main.cell(row, 1).value = idx + 1 if idx < ranked else "미산출"
        main.cell(row, 3).value = code
        main.cell(row, 4).value = f'=IF(ISNUMBER(원자료!I{raw_row}),원자료!I{raw_row},"미확인")'
        main.cell(row, 5).value = f'=IF(AND(ISNUMBER(D{row}),ISNUMBER(원자료!J{raw_row})),D{row}*원자료!J{raw_row}/100000000,"미산출")'
        for col, raw_col in ((6, "T"), (7, "D"), (9, "U"), (10, "E"), (14, "O")):
            main.cell(row, col).value = f'=IF(ISNUMBER(원자료!{raw_col}{raw_row}),원자료!{raw_col}{raw_row},"미확인")'
        main.cell(row, 8).value = yoy_formula("F", "G", row)
        main.cell(row, 11).value = yoy_formula("I", "J", row)
        main.cell(row, 12).value = f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(G{row}),G{row}>0),E{row}/G{row},"미산출")'
        main.cell(row, 13).value = f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(J{row}),J{row}>0),E{row}/J{row},"미산출")'
        values = {"D": quotes[code].price, "E": cap[code], "F": profits[code][0],
                  "G": profits[code][1], "H": yoy(profits[code][0], profits[code][1]),
                  "I": profits[code][2], "J": profits[code][3],
                  "K": yoy(profits[code][2], profits[code][3]),
                  "L": per[code], "M": por[code], "N": profits[code][4]}
        for col, value in values.items():
            cache[f"{col}{row}"] = value if isinstance(value, (int, float)) else (
                value if isinstance(value, str) else "미확인" if col in ("D", "F", "G", "H", "I", "J", "K", "N") else "미산출")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    _set_cached_values(output, cache)
    validate(template, output, quotes, per, por)
    return {"successful_prices": sum(q.price is not None and q.shares is not None for q in quotes.values()),
            "missing": [code for code in ordered if quotes[code].price is None or quotes[code].shares is None],
            "uncomputed_per": [code for code in ordered if per[code] is None]}


def validate(template, output, quotes, expected_per, expected_por):
    original, old_rows = records(template)
    current = load_workbook(output, data_only=False)
    cached = load_workbook(output, data_only=True)
    if current.sheetnames != original.sheetnames or current[MAIN].tables["ValuationTable"].ref != original[MAIN].tables["ValuationTable"].ref:
        raise ValueError("output sheet/table structure changed")
    raw = current[RAW]
    main = current[MAIN]
    if raw.tables["SourceTable"].ref != "A9:AB99":
        raise ValueError("source provenance table structure changed")
    if {str(raw.cell(r, 2).value) for r in range(FIRST, LAST + 1)} != set(old_rows):
        raise ValueError("stock mapping changed")
    main_codes = [str(main.cell(r, 3).value) for r in range(FIRST, LAST + 1)]
    if len(set(main_codes)) != 90 or set(main_codes) != set(old_rows):
        raise ValueError("output stock loss or duplication")
    for code, old_row in old_rows.items():
        for col in range(1, 25):
            if col in (9, 10, 11, 13, 14):
                continue
            a = raw.cell(old_row, col).value
            b = original[RAW].cell(old_row, col).value
            same = isclose(a, b, rel_tol=0, abs_tol=1e-12) if isinstance(a, (int, float)) and isinstance(b, (int, float)) else a == b
            if not same:
                raise ValueError("source estimate, ROE, or note changed")
        row = FIRST + main_codes.index(code)
        q = quotes[code]
        if cached[MAIN].cell(row, 4).value != (q.price if q.price is not None else "미확인"):
            raise ValueError("cached price mismatch")
        cap = q.price * q.shares / 100_000_000 if q.price is not None and q.shares is not None else None
        source_row = original[RAW]
        derived = {
            5: cap, 6: source_row.cell(old_row, 20).value,
            7: source_row.cell(old_row, 4).value,
            8: yoy(source_row.cell(old_row, 20).value, source_row.cell(old_row, 4).value),
            9: source_row.cell(old_row, 21).value,
            10: source_row.cell(old_row, 5).value,
            11: yoy(source_row.cell(old_row, 21).value, source_row.cell(old_row, 5).value),
            12: expected_per[code], 13: expected_por[code],
            14: source_row.cell(old_row, 15).value,
        }
        for col, expected in derived.items():
            actual = cached[MAIN].cell(row, col).value
            if expected is None:
                if actual != ("미산출" if col in (5, 12, 13) else "미확인"):
                    raise ValueError("missing calculation was filled")
            elif isinstance(expected, (int, float)) and (not isinstance(actual, (int, float)) or abs(actual - expected) > 1e-8):
                raise ValueError("valuation calculation mismatch")
            elif isinstance(expected, str) and actual != expected:
                raise ValueError("YoY calculation mismatch")
        if not all(main.cell(row, col).data_type == "f" for col in range(4, 15)):
            raise ValueError("valuation formula missing")
        old_main_row = next(r for r in range(FIRST, LAST + 1) if str(original[MAIN].cell(r, 3).value) == code)
        for col in (2, 15, 16, 17):
            if main.cell(row, col).value != original[MAIN].cell(old_main_row, col).value:
                raise ValueError("stock detail shifted after sorting")
        if (raw.cell(old_row, 28).value or "") != (q.reason or ""):
            raise ValueError("source exception reason mismatch")
    sorted_per = [expected_per[c] for c in main_codes if expected_per[c] is not None]
    if sorted_per != sorted(sorted_per):
        raise ValueError("PER sort order incorrect")
