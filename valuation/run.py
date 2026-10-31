"""One dated valuation run; safe to invoke from workflow_dispatch or schedule."""
from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date, datetime, time as clock_time, timedelta
from pathlib import Path

from openpyxl import load_workbook

from .api import DataError, KRXCloseEmpty, KST, MarketData, Quote
from .delivery import DeliveryError, DuplicateDelivery, Ledger, Telegram
from .workbook import MAIN, RAW, build, records

TEMPLATE = Path(__file__).parent / "template" / "2027_valuation_20261020_133.xlsx"
CLOSE_CRON = "47 10 * * 2-6"


def flag(value):
    return str(value).lower() in ("true", "1", "yes")


def latest_scheduled_day(now, slot):
    """Resolve the preceding day's close for the last Tue-Sat 10:47 run."""
    if slot != "close":
        raise ValueError("only close has a scheduled target")
    trigger = clock_time(10, 47)
    candidate = now.date()
    if now.time() < trigger:
        candidate -= timedelta(days=1)
    while candidate.weekday() not in (1, 2, 3, 4, 5):
        candidate -= timedelta(days=1)
    return candidate - timedelta(days=1)


def scheduled_target(now, cron):
    if cron == CLOSE_CRON:
        return "close", latest_scheduled_day(now, "close")
    raise ValueError("unrecognized valuation schedule")


def argument_parser():
    parser = argparse.ArgumentParser()
    now = datetime.now(KST)
    if os.getenv("GITHUB_EVENT_NAME") == "schedule":
        slot, default_day = scheduled_target(now, os.getenv("SCHEDULE_CRON", ""))
    else:
        slot, default_day = os.getenv("SLOT", ""), now.date()
    parser.add_argument("--trade-date", default=os.getenv("TRADE_DATE") or default_day.isoformat())
    parser.add_argument("--slot", choices=("close",), default=slot or "close")
    parser.add_argument("--dry-run", action="store_true", default=flag(os.getenv("DRY_RUN", "false")))
    parser.add_argument("--test-send", action="store_true", default=flag(os.getenv("TEST_SEND", "false")))
    parser.add_argument("--resend", action="store_true", default=flag(os.getenv("RESEND", "false")))
    parser.add_argument("--fixture-template-prices", action="store_true")
    parser.add_argument("--output-dir", default=os.getenv("OUTPUT_DIR", "valuation-output"))
    return parser


def enforce_time(day, slot, now):
    if day > now.date():
        raise ValueError("future trade date is not allowed")
    if day == now.date():
        threshold = clock_time(11, 35) if slot == "1130" else clock_time(15, 2) if slot == "1500" else clock_time(16, 0)
        if now.time() < threshold:
            raise ValueError("price basis time has not been reached")


def fixture_quotes(workbook, rows):
    raw = workbook["원자료"]
    stamp = datetime.now(KST).isoformat(timespec="seconds")
    return {code: Quote(raw.cell(row, 9).value, raw.cell(row, 10).value,
                        "테스트: 첨부 템플릿 기존 값 (실시간 검증 아님)",
                        "테스트 자료", stamp) for code, row in rows.items()}


def caption(day, slot, summary, test=False):
    missing = summary["missing"]
    suffix = ("없음" if not missing else ", ".join(missing[:12]) +
              (f" 외 {len(missing)-12}종목" if len(missing) > 12 else ""))
    return (f"{'[테스트] ' if test else ''}2027 밸류에이션 {day.isoformat()}\n"
            f"가격 기준: {'11:30 KST 또는 그 이전 KRX 정규장 마지막 분봉(개별 실제시각은 원자료)' if slot == '1130' else '15:00 KST 또는 그 이전 KRX 정규장 마지막 분봉(개별 실제시각은 원자료)' if slot == '1500' else '15:30 KST KRX 정규장 종가'}\n"
            f"가격·주식수 확인: {summary['successful_prices']}/{summary['total_stocks']}종목\n"
            f"27E PER/POR 미산출: {len(summary['uncomputed_per'])}/{len(summary['uncomputed_por'])}종목\n"
            f"{'부분 누락 있음 — ' if missing else ''}가격·주식수 미확인: {suffix}")


def write_log(folder, name, data):
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def verify_close_quotes(quotes, codes, day):
    if set(quotes) != set(codes):
        raise DataError("close verification failed: stock universe mismatch")
    for q in quotes.values():
        if (q.trade_date != day.isoformat() or q.basis != "15:30 KST 정규장 종가"
                or q.source not in ("KRX stk_bydd_trd", "KRX ksq_bydd_trd")
                or q.price is None or q.price <= 0 or q.shares is None or q.shares <= 0):
            raise DataError("close verification failed: date, session, source or value mismatch")


def verify_close_file(path, rows, quotes, day):
    workbook = load_workbook(path, data_only=True)
    try:
        raw = workbook[RAW]
        if (not str(workbook[MAIN]["A3"].value).startswith(day.isoformat())
                or not str(raw["A4"].value).startswith(day.isoformat())):
            raise DataError("workbook trade date mismatch")
        for code, row in rows.items():
            q = quotes[code]
            if (str(raw.cell(row, 2).value) != code
                    or raw.cell(row, 9).value != q.price
                    or raw.cell(row, 10).value != q.shares
                    or raw.cell(row, 25).value != q.source
                    or raw.cell(row, 26).value != q.basis):
                raise DataError("workbook close evidence mismatch")
    finally:
        workbook.close()


def verify_send_day(day, *, allow_historical=False):
    now = datetime.now(KST)
    if day > now.date():
        raise ValueError("live delivery cannot use a future KST trade date")
    if day < now.date() and not allow_historical:
        raise ValueError("historical delivery requires an explicit manual trade date or scheduled close run")
    if day == now.date():
        enforce_time(day, "close", now)


def main(argv=None):
    args = argument_parser().parse_args(argv)
    folder = Path(args.output_dir)
    schedule_cron = os.getenv("SCHEDULE_CRON", "") if os.getenv("GITHUB_EVENT_NAME") == "schedule" else ""
    log = {"trade_date": args.trade_date, "slot": args.slot,
           "dry_run": args.dry_run, "test_send": args.test_send,
           "fixture": args.fixture_template_prices,
           "started_at": datetime.now(KST).isoformat(timespec="seconds")}
    if schedule_cron:
        log["schedule_cron"] = schedule_cron
    claimed = False
    ledger = None
    try:
        if args.slot != "close":
            raise ValueError("only the KRX regular-session close slot is allowed")
        if args.fixture_template_prices and not args.dry_run:
            raise ValueError("fixture prices are allowed only in dry-run")
        if args.resend and (args.dry_run or os.getenv("GITHUB_EVENT_NAME") != "workflow_dispatch"):
            raise ValueError("resend requires an explicit manual send run")
        day = date.fromisoformat(args.trade_date)
        now = datetime.now(KST)
        if schedule_cron and day != scheduled_target(now, schedule_cron)[1]:
            log["status"] = "stale_schedule_skipped"
            write_log(folder, "run.json", log)
            print(json.dumps(log, ensure_ascii=False))
            return 0
        allow_historical = bool(schedule_cron) or (
            os.getenv("GITHUB_EVENT_NAME") == "workflow_dispatch"
            and bool(os.getenv("TRADE_DATE", "").strip()))
        if not args.dry_run:
            verify_send_day(day, allow_historical=allow_historical)
        if not args.fixture_template_prices:
            enforce_time(day, args.slot, now)
        key = f"{day:%Y%m%d}_{args.slot}"
        log["ledger_key"] = key
        if not args.dry_run and not args.resend:
            ledger = Ledger()
            previous, _ = ledger.read(key)
            if previous:
                raise DuplicateDelivery(previous["status"])
        template, rows = records(TEMPLATE)
        if args.fixture_template_prices:
            quotes = fixture_quotes(template, rows)
        else:
            market = MarketData()
            try:
                quotes = market.collect_verified_close(rows, day)
            except KRXCloseEmpty:
                # A complete KRX dataset proves this was an open session. Only
                # an empty dataset needs the separate KIS calendar request.
                if market.is_open(day):
                    raise
                if not args.dry_run:
                    ledger = ledger or Ledger()
                    ledger.ensure_branch()
                    previous, sha = ledger.read(key)
                    if not previous:
                        ledger.write(key, {"status": "market_closed", "trade_date": day.isoformat()}, sha)
                log["status"] = "market_closed"
                write_log(folder, "run.json", log)
                print("KRX market closed; skipped")
                return 0
            verify_close_quotes(quotes, rows, day)
            log["price_verification"] = {"trade_date": day.isoformat(),
                                         "basis": "15:30 KST 정규장 종가",
                                         "verified_stocks": len(quotes), "passed": True}
        out = folder / f"valuation_{day:%Y%m%d}_{args.slot}.xlsx"
        summary = build(TEMPLATE, out, day, args.slot, quotes,
                        test=args.fixture_template_prices)
        log.update({"status": "validated", "file": out.name, **summary,
                    "checked_at": datetime.now(KST).isoformat(timespec="seconds")})
        write_log(folder, "run.json", log)
        if args.dry_run:
            print(json.dumps(log, ensure_ascii=False))
            return 0
        verify_close_quotes(quotes, rows, day)
        verify_close_file(out, rows, quotes, day)
        verify_send_day(day, allow_historical=allow_historical)
        telegram = Telegram()
        telegram.check_permission()
        ledger = ledger or Ledger()
        claim = ledger.claim(key, force=args.resend)
        claimed = True
        verify_send_day(day, allow_historical=allow_historical)
        message_id = telegram.send_document(out, caption(day, args.slot, summary, args.test_send))
        ledger.complete(key, claim, message_id)
        log.update({"status": "sent", "telegram_message_id": message_id})
        write_log(folder, "run.json", log)
        print(json.dumps(log, ensure_ascii=False))
        return 0
    except DuplicateDelivery as exc:
        log.update({"status": "duplicate_blocked", "ledger_status": exc.status})
        write_log(folder, "run.json", log)
        print(json.dumps(log, ensure_ascii=False))
        return 0 if exc.status in ("sent", "market_closed") else 1
    except (ValueError, DataError, DeliveryError, OSError) as exc:
        # Error classes/messages are intentionally narrow. Do not log request URLs or tokens.
        reason = str(exc)
        log.update({"status": "failed", "error": reason})
        write_log(folder, "run.json", log)
        if not args.dry_run and not args.fixture_template_prices:
            try:
                Telegram().send_failure(args.trade_date, args.slot, reason, uncertain=claimed)
            except DeliveryError:
                pass
        print(json.dumps(log, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
