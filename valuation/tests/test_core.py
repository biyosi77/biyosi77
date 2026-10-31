from datetime import date, datetime
import json
from math import isclose
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
import requests
from openpyxl import load_workbook

from valuation.api import DataError, KIS, KRXCloseEmpty, MarketData, Quote, choose_minute, credential
from valuation.delivery import DeliveryError, Ledger
from valuation.run import caption, enforce_time, latest_scheduled_day, scheduled_target
from valuation.workbook import MAIN, RAW, build, records, yoy, yoy_formula


TEMPLATE = Path(__file__).parents[1] / "template" / "2027_valuation_20261020_133.xlsx"
BASE_TEMPLATE = Path(__file__).parents[1] / "template" / "2027_valuation_20261016_yoy.xlsx"
ADDITIONS = Path(__file__).parents[1] / "template" / "additions_20261020.json"
DAY = date(2026, 10, 15)


def test_price_time_selects_latest_prior_bar_on_requested_day():
    bars = [
        {"stck_bsop_date": "20261014", "stck_cntg_hour": "113000", "stck_prpr": "999"},
        {"stck_bsop_date": "20261015", "stck_cntg_hour": "113100", "stck_prpr": "888"},
        {"stck_bsop_date": "20261015", "stck_cntg_hour": "112900", "stck_prpr": "101"},
        {"stck_bsop_date": "20261015", "stck_cntg_hour": "113000", "stck_prpr": "102"},
    ]
    assert choose_minute(bars, DAY) == ("113000", 102)
    assert choose_minute(bars[:-1], DAY) == ("112900", 101)
    assert choose_minute(bars[:2], DAY) is None


def test_1500_price_uses_exact_bar_and_rejects_later_or_previous_day():
    bars = [
        {"stck_bsop_date": "20261014", "stck_cntg_hour": "150000", "stck_prpr": "999"},
        {"stck_bsop_date": "20261015", "stck_cntg_hour": "145900", "stck_prpr": "101"},
        {"stck_bsop_date": "20261015", "stck_cntg_hour": "150000", "stck_prpr": "102"},
        {"stck_bsop_date": "20261015", "stck_cntg_hour": "150100", "stck_prpr": "888"},
    ]
    assert choose_minute(bars, DAY, "150000") == ("150000", 102)
    assert choose_minute(bars[:2], DAY, "150000") == ("145900", 101)


def test_invalid_api_secret_is_not_retried_or_logged():
    class Session:
        calls = 0
        def request(self, *args, **kwargs):
            self.calls += 1
            response = requests.Response()
            response.status_code = 403
            response.headers["Content-Type"] = "application/json"
            response._content = b'{"error_code":"EGW00105","appsecret":"sensitive"}'
            return response
    client = MarketData.__new__(MarketData)
    client.session = Session()
    client.sleep = lambda _: None
    with pytest.raises(DataError, match=r"KIS token request failed \(HTTP 403, code EGW00105\)") as error:
        client._request("POST", KIS + "/oauth2/tokenP", json={})
    assert client.session.calls == 1
    assert "sensitive" not in str(error.value)


def test_kis_token_rate_limit_waits_once_then_retries():
    class Session:
        calls = 0
        def request(self, *args, **kwargs):
            self.calls += 1
            response = requests.Response()
            response.headers["Content-Type"] = "application/json"
            if self.calls == 1:
                response.status_code = 403
                response._content = b'{"error_code":"EGW00133"}'
            else:
                response.status_code = 200
                response._content = b'{"access_token":"ok"}'
            return response
    client = MarketData.__new__(MarketData)
    client.session = Session()
    waits = []
    client.sleep = waits.append
    assert client._request("POST", KIS + "/oauth2/tokenP") == {"access_token": "ok"}
    assert client.session.calls == 2
    assert waits == [65]


def test_secret_copy_format_is_trimmed(monkeypatch):
    monkeypatch.setenv("KIS_APP_SECRET", '  "abc123"  ')
    assert credential("KIS_APP_SECRET") == "abc123"


def test_holiday_requires_exact_day_and_open_flag():
    client = MarketData.__new__(MarketData)
    client.kis = lambda *a: {"output": [{"bass_dt": "20261015", "opnd_yn": "N"}]}
    assert client.is_open(DAY) is False
    client.kis = lambda *a: {"output": [{"bass_dt": "20261014", "opnd_yn": "Y"}]}
    with pytest.raises(Exception, match="requested date"):
        client.is_open(DAY)


def test_minute_request_is_krx_venue_and_requested_day():
    client = MarketData.__new__(MarketData)
    calls = []
    def fake_kis(path, tr_id, params):
        calls.append((path, tr_id, params))
        return {"output2": [{"stck_bsop_date": "20261015", "stck_cntg_hour": "113000", "stck_prpr": "100"}]}
    client.kis = fake_kis
    assert client.minute("005930", DAY) == ("113000", 100)
    assert calls[0][2]["FID_COND_MRKT_DIV_CODE"] == "J"
    assert calls[0][2]["FID_INPUT_DATE_1"] == "20261015"
    assert calls[0][2]["FID_INPUT_HOUR_1"] == "113000"
    client.minute("005930", DAY, "150000")
    assert calls[-1][2]["FID_INPUT_HOUR_1"] == "150000"


def test_1500_same_day_uses_live_krx_minute_endpoint():
    client = MarketData.__new__(MarketData)
    calls = []
    today = datetime.now(ZoneInfo("Asia/Seoul")).date()
    def fake_kis(path, tr_id, params):
        calls.append((path, tr_id, params))
        return {"output2": [{"stck_bsop_date": today.strftime("%Y%m%d"),
                              "stck_cntg_hour": "150000", "stck_prpr": "100"}]}
    client.kis = fake_kis
    assert client.minute("005930", today, "150000") == ("150000", 100)
    assert calls[0][0].endswith("/inquire-time-itemchartprice")
    assert calls[0][1] == "FHKST03010200"
    assert calls[0][2]["FID_COND_MRKT_DIV_CODE"] == "J"
    client.kis = lambda path, tr_id, params: (
        calls.append((path, tr_id, params)) or
        {"output2": [{"stck_bsop_date": today.strftime("%Y%m%d"),
                      "stck_cntg_hour": "153000", "stck_prpr": "225000"},
                     {"stck_bsop_date": today.strftime("%Y%m%d"),
                      "stck_cntg_hour": "160000", "stck_prpr": "220500"}]})
    assert client.minute("005930", today, "153000") == ("153000", 225000)
    assert calls[-1][0].endswith("/inquire-time-itemchartprice")
    assert calls[-1][2]["FID_INPUT_HOUR_1"] == "153000"


def test_close_rejects_wrong_date():
    client = MarketData.__new__(MarketData)
    client.krx = lambda *a: [{"ISU_CD": "005930", "BAS_DD": "20261014",
                              "TDD_CLSPRC": "100", "LIST_SHRS": "1000"}]
    with pytest.raises(Exception, match="empty"):
        client.closes(DAY)


def test_legacy_collector_cannot_use_after_market_daily_price_for_close():
    client = MarketData.__new__(MarketData)
    with pytest.raises(DataError, match="invalid price slot"):
        client.collect(["005930"], date(2026, 10, 19), "close")


def test_reference_time_guard():
    kst = ZoneInfo("Asia/Seoul")
    with pytest.raises(ValueError):
        enforce_time(DAY, "close", datetime(2026, 10, 15, 15, 39, tzinfo=kst))
    enforce_time(DAY, "1130", datetime(2026, 10, 15, 12, 0, tzinfo=kst))
    assert latest_scheduled_day(datetime(2026, 10, 30, 15, 47, tzinfo=kst), "close") == date(2026, 10, 30)
    assert latest_scheduled_day(datetime(2026, 10, 31, 0, 30, tzinfo=kst), "close") == date(2026, 10, 30)
    with pytest.raises(ValueError, match="only close"):
        latest_scheduled_day(datetime(2026, 10, 31, 10, 47, tzinfo=kst), "1130")


def test_single_schedule_targets_same_day_close():
    kst = ZoneInfo("Asia/Seoul")
    cron = "47 15 * * 1-5"
    assert scheduled_target(datetime(2026, 10, 30, 15, 47, tzinfo=kst), cron) == ("close", date(2026, 10, 30))
    assert scheduled_target(datetime(2026, 10, 31, 0, 30, tzinfo=kst), cron) == ("close", date(2026, 10, 30))
    assert scheduled_target(datetime(2026, 11, 2, 16, 0, tzinfo=kst), cron) == ("close", date(2026, 11, 2))
    with pytest.raises(ValueError, match="unrecognized"):
        scheduled_target(datetime(2026, 10, 31, 10, 47, tzinfo=kst), "17 16 * * 1-5")


def test_schedule_parser_keeps_resolved_close_with_empty_env(monkeypatch):
    import valuation.run as run
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 10, 30, 15, 50, tzinfo=tz)
    monkeypatch.setattr(run, "datetime", FrozenDateTime)
    monkeypatch.setenv("GITHUB_EVENT_NAME", "schedule")
    monkeypatch.setenv("SCHEDULE_CRON", "47 15 * * 1-5")
    monkeypatch.setenv("SLOT", "")
    monkeypatch.setenv("TRADE_DATE", "")
    args = run.argument_parser().parse_args([])
    assert (args.trade_date, args.slot) == ("2026-10-30", "close")


def test_manual_nonclose_rejected_before_delivery(monkeypatch, tmp_path):
    import valuation.run as run
    monkeypatch.setenv("GITHUB_EVENT_NAME", "workflow_dispatch")
    monkeypatch.setenv("SLOT", "1130")
    monkeypatch.setenv("DRY_RUN", "true")
    assert run.main(["--output-dir", str(tmp_path)]) == 1
    assert json.loads((tmp_path / "run.json").read_text())["error"] == "only the KRX regular-session close slot is allowed"


def test_empty_price_slot_cannot_fall_back_to_minute_data():
    client = MarketData.__new__(MarketData)
    with pytest.raises(DataError, match="invalid price slot"):
        client.collect(["005930"], DAY, "")


def test_yoy_calculation_matches_new_template():
    assert yoy(100, 125) == pytest.approx(0.25)
    assert yoy(100, 0) == -1
    assert yoy(100, -1) == "적자전환"
    assert yoy(-100, 10) == "흑자전환"
    assert yoy(-100, -50) == "적자지속"
    assert yoy(0, 10) == "흑자전환"
    assert yoy(0, 0) == "손익분기"
    assert yoy(0, -10) == "미확인"
    assert yoy(None, 10) == "미확인"
    assert "F10" in yoy_formula("F", "G", 10)
    assert "G10/F10-1" in yoy_formula("F", "G", 10)


def test_calculation_code_mapping_sort_and_missing(tmp_path):
    original, rows = records(TEMPLATE)
    added = json.loads(ADDITIONS.read_text(encoding="utf-8"))["stocks"]
    base, base_rows = records(BASE_TEMPLATE)
    assert len(rows) == 133 and len(base_rows) == 90 and len(added) == 43
    assert set(base_rows).issubset(rows)
    assert {item["code"] for item in added} == set(rows) - set(base_rows)
    for code, base_row in base_rows.items():
        for col in range(1, 25):
            actual = original[RAW].cell(rows[code], col).value
            expected = base[RAW].cell(base_row, col).value
            if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
                assert isclose(actual, expected, rel_tol=0, abs_tol=1e-12)
            else:
                assert actual == expected
    for item in added:
        row = rows[item["code"]]
        assert original[RAW].cell(row, 2).value == item["code"]
        assert original[RAW].cell(row, 1).value == item["name"]
        assert all(original[RAW].cell(row, col).value is None
                   for col in (4, 5, 15, 20, 21))
    assert original[MAIN].tables["ValuationTable"].ref == "A9:Q142"
    source = original[RAW]
    quotes = {code: Quote(source.cell(row, 9).value, source.cell(row, 10).value,
                          "fixture KRX", "11:30 KST", "2026-10-15T12:01:00+09:00")
              for code, row in rows.items()}
    missing = "005930"
    quotes[missing] = Quote(None, None, "fixture KRX", "", "2026-10-15T12:01:00+09:00", "no bar")
    added_code = "078350"
    quotes[added_code] = Quote(10_000, 20_000_000, "fixture KRX", "11:30 KST",
                               "2026-10-15T12:01:00+09:00")
    output = tmp_path / "valuation_20261015_1130.xlsx"
    summary = build(TEMPLATE, output, DAY, "1130", quotes, test=True)
    formulas = load_workbook(output, data_only=False)
    values = load_workbook(output, data_only=True)
    assert formulas.sheetnames == original.sheetnames
    assert formulas[MAIN]["N9"].value == "27E ROE"
    assert formulas[MAIN]["D9"].value == "10/15 11:30 주가"
    assert formulas[RAW]["I9"].value == "10/15 11:30 주가"
    assert formulas[RAW]["AB9"].value == "예외 사유"
    assert formulas[RAW]["T9"].value == original[RAW]["T9"].value
    assert formulas[RAW].tables["SourceTable"].ref == "A9:AB142"
    assert summary["total_stocks"] == 133
    assert "133종목" in caption(DAY, "1130", summary)
    assert "27E PER/POR 미산출:" in caption(DAY, "1130", summary)
    assert added_code in summary["uncomputed_per"]
    assert missing in summary["missing"]
    ordered = [formulas[MAIN].cell(r, 3).value for r in range(10, 143)]
    row = 10 + ordered.index(missing)
    assert values[MAIN].cell(row, 1).value == "미산출"
    assert values[MAIN].cell(row, 4).value == "미확인"
    assert values[MAIN].cell(row, 5).value == "미산출"
    assert values[MAIN].cell(row, 12).value == "미산출"
    assert values[RAW].cell(rows[missing], 28).value == "no bar"
    added_row = 10 + ordered.index(added_code)
    assert values[MAIN].cell(added_row, 5).value == 2000
    assert values[MAIN].cell(added_row, 7).value == "미확인"
    assert values[MAIN].cell(added_row, 12).value == "미산출"
    assert values[MAIN].cell(added_row, 17).value == "추가 / 추정치 미제공"
    for code in rows:
        raw_row = rows[code]
        main_row = 10 + ordered.index(code)
        assert values[MAIN].cell(main_row, 6).value == source.cell(raw_row, 20).value or (
            source.cell(raw_row, 20).value is None and values[MAIN].cell(main_row, 6).value == "미확인")
        assert values[MAIN].cell(main_row, 7).value == source.cell(raw_row, 4).value or (
            source.cell(raw_row, 4).value is None and values[MAIN].cell(main_row, 7).value == "미확인")
        assert values[MAIN].cell(main_row, 14).value == source.cell(raw_row, 15).value or (
            source.cell(raw_row, 15).value is None and values[MAIN].cell(main_row, 14).value == "미확인")
        assert values[RAW].cell(raw_row, 20).value == source.cell(raw_row, 20).value
        assert values[RAW].cell(raw_row, 24).value == source.cell(raw_row, 24).value
        assert formulas[MAIN].cell(main_row, 17).value == next(
            original[MAIN].cell(r, 17).value for r in range(10, 143)
            if original[MAIN].cell(r, 3).value == code)


def test_delivery_claim_blocks_rerun_and_marks_after_success():
    class MemoryLedger(Ledger):
        def __init__(self):
            self.value = None
        def ensure_branch(self):
            pass
        def read(self, key):
            return self.value, "sha" if self.value else None
        def write(self, key, value, sha=None):
            self.value = value
    ledger = MemoryLedger()
    claim = ledger.claim("20261015_1130")
    assert ledger.value["status"] == "sending"
    with pytest.raises(DeliveryError):
        ledger.claim("20261015_1130")
    ledger.complete("20261015_1130", claim, 123)
    assert ledger.value["status"] == "sent"
    assert ledger.value["telegram_message_id"] == 123
    with pytest.raises(DeliveryError):
        ledger.claim("20261015_1130")
    ledger.claim("20261015_1130", force=True)
    assert ledger.value["status"] == "sending"


def test_close_delivery_uses_separate_key_and_blocks_retry(monkeypatch, tmp_path):
    import valuation.run as run
    values = {"20261029_": {"status": "sent", "telegram_message_id": 173}}
    sends = []
    market_calls = []
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None): return cls(2026, 10, 29, 18, 0, tzinfo=tz)
    monkeypatch.setattr(run, "datetime", FrozenDateTime)
    class FakeLedger(Ledger):
        def __init__(self): pass
        def ensure_branch(self): pass
        def read(self, key): return values.get(key), "sha"
        def write(self, key, value, sha=None): values[key] = value
    class FakeMarket:
        def is_open(self, day): return True
        def collect_same_day_close(self, rows, day): return self.collect_verified_close(rows, day)
        def collect_verified_close(self, rows, day):
            market_calls.append(day)
            return {code: Quote(10000, 1000000, "KRX stk_bydd_trd", "15:30 KST 정규장 종가", "test", trade_date=day.isoformat()) for code in rows}
    class FakeTelegram:
        def check_permission(self): pass
        def send_document(self, path, text):
            sends.append((path.name, text)); return 999
    monkeypatch.setattr(run, "Ledger", FakeLedger)
    monkeypatch.setattr(run, "MarketData", FakeMarket)
    monkeypatch.setattr(run, "Telegram", FakeTelegram)
    monkeypatch.setenv("GITHUB_EVENT_NAME", "workflow_dispatch")
    monkeypatch.setenv("SLOT", "close")
    monkeypatch.setenv("DRY_RUN", "false")
    monkeypatch.setenv("RESEND", "false")
    args = ["--trade-date", "2026-10-29", "--output-dir", str(tmp_path)]
    assert run.main(args) == 0
    assert values["20261029_close"]["status"] == "sent"
    _, rows = records(TEMPLATE)
    quotes = FakeMarket().collect_verified_close(rows, date(2026, 10, 29))
    market_calls.pop()  # verification fixture is not a production collection
    output = tmp_path / "valuation_20261029_close.xlsx"
    workbook = load_workbook(output)
    workbook[RAW]["A4"] = "2026-10-28 stale date"
    workbook.save(output)
    with pytest.raises(DataError, match="trade date mismatch"):
        run.verify_close_file(output, rows, quotes, date(2026, 10, 29))
    assert run.main(args) == 0
    assert len(sends) == 1
    assert len(market_calls) == 1
    log = json.loads((tmp_path / "run.json").read_text())
    assert log["status"] == "duplicate_blocked"
    assert log["ledger_key"] == "20261029_close"


@pytest.mark.parametrize('field,value', [('trade_date','2026-10-28'), ('basis','11:30 KST'), ('price',None), ('source','KIS KRX 분봉')])
def test_close_verification_rejects_unverified_quotes(field, value):
    from dataclasses import replace
    from valuation.run import verify_close_quotes
    day = date(2026, 10, 29)
    q = Quote(100, 1000, 'KRX stk_bydd_trd', '15:30 KST 정규장 종가', 'test', trade_date=day.isoformat())
    with pytest.raises(DataError, match='verification failed'):
        verify_close_quotes({'005930': replace(q, **{field: value})}, ['005930'], day)


def test_verified_close_is_two_bulk_requests_and_checks_response_date():
    client = MarketData.__new__(MarketData)
    calls = []
    def krx(endpoint, day):
        calls.append(endpoint)
        return [{'ISU_SRT_CD': '005930', 'BAS_DD': '20261029', 'TDD_CLSPRC': '100', 'LIST_SHRS': '1000'}] if endpoint == 'stk_bydd_trd' else []
    client.krx = krx
    quotes = client.collect_verified_close(['005930'], date(2026, 10, 29))
    assert quotes['005930'].trade_date == '2026-10-29'
    assert calls == ['stk_bydd_trd', 'ksq_bydd_trd']
    with pytest.raises(DataError, match='empty or not yet published'):
        client.collect_verified_close(['005930'], date(2026, 10, 30))
    with pytest.raises(DataError, match='incomplete'):
        client.collect_verified_close(['005930','000660'], date(2026, 10, 29))


def test_kis_regular_close_uses_exact_1530_bar_and_separate_shares():
    client = MarketData.__new__(MarketData)
    calls = []
    def minute(code, day, cutoff):
        calls.append(('minute', code, day, cutoff))
        return ('153000', 225000)
    def shares(code):
        calls.append(('shares', code))
        return 5846278608
    client.minute = minute
    client.current_shares = shares
    quote = client.kis_regular_close_snapshot('005930', date(2026, 10, 30))
    assert (quote.price, quote.shares, quote.trade_date) == (225000, 5846278608, '2026-10-30')
    assert quote.basis == '15:30 KST 정규장 종가'
    assert calls == [('minute', '005930', date(2026, 10, 30), '153000'),
                     ('shares', '005930')]
    calls.clear()
    client.minute = lambda *args: ('152900', 220500)
    stale = client.kis_regular_close_snapshot('005930', date(2026, 10, 30))
    assert stale.price is None and '15:30' in stale.reason and calls == []


def test_same_day_collection_polls_one_sentinel_then_fetches_each_code_once():
    client = MarketData.__new__(MarketData)
    calls = []
    sleeps = []
    def snapshot(code, day):
        calls.append(code)
        price = None if len(calls) == 1 else 100
        return Quote(price, 1000, 'KIS KRX(J) 15:30 분봉; KIS 상장주식수',
                     '15:30 KST 정규장 종가' if price else '', 'test',
                     '' if price else '당일 종가 미확인', day.isoformat())
    client.kis_regular_close_snapshot = snapshot
    client.sleep = sleeps.append
    quotes = client.collect_same_day_close(['005930', '000660'], date(2026, 10, 30),
        now=lambda: datetime(2026, 10, 30, 16, 0, tzinfo=ZoneInfo('Asia/Seoul')))
    assert set(quotes) == {'005930', '000660'}
    assert calls == ['005930', '005930', '000660']
    assert sleeps == [300, 0.2]


def test_same_day_partial_missing_is_marked_without_stale_substitution(tmp_path):
    from valuation.run import verify_close_file, verify_close_quotes
    client = MarketData.__new__(MarketData)
    client.sleep = lambda _: None
    def snapshot(code, day):
        price = 100 if code == '005930' else None
        return Quote(price, 1000, 'KIS KRX(J) 15:30 분봉; KIS 상장주식수',
                     '15:30 KST 정규장 종가' if price else '', 'test',
                     '' if price else '당일 종가 미확인', day.isoformat())
    client.kis_regular_close_snapshot = snapshot
    day = date(2026, 10, 30)
    quotes = client.collect_same_day_close(['005930', '000660'], day,
        now=lambda: datetime(2026, 10, 30, 16, 45, tzinfo=ZoneInfo('Asia/Seoul')))
    verify_close_quotes(quotes, ['005930', '000660'], day)
    assert quotes['000660'].price is None and quotes['000660'].reason
    _, rows = records(TEMPLATE)
    full = {code: Quote(100, 1000, 'KIS KRX(J) 15:30 분봉; KIS 상장주식수',
                        '15:30 KST 정규장 종가', 'test', trade_date=day.isoformat())
            for code in rows}
    full['172670'] = Quote(None, None, 'KIS KRX(J) 15:30 분봉; KIS 상장주식수',
                           '', 'test', '15:30 KRX 정규장 분봉 미확인', day.isoformat())
    output = tmp_path / 'valuation_20261030_close.xlsx'
    summary = build(TEMPLATE, output, day, 'close', full)
    verify_close_quotes(full, rows, day)
    verify_close_file(output, rows, full, day)
    assert summary['successful_prices'] == 132
    assert summary['missing'] == ['172670']
    assert '부분 누락 있음' in caption(day, 'close', summary, delayed=True)


def test_rate_limit_stops_without_request_retry():
    class Session:
        calls = 0
        def request(self, *args, **kwargs):
            self.calls += 1
            response = requests.Response(); response.status_code = 429
            return response
    client = MarketData.__new__(MarketData)
    client.session = Session()
    client.sleep = lambda _: pytest.fail('must not retry rate limit')
    with pytest.raises(DataError, match='rate limited'):
        client._request('GET', 'https://data-dbg.krx.co.kr/test')
    assert client.session.calls == 1


def test_kis_session_error_is_retried_only_twice():
    client = MarketData.__new__(MarketData)
    client.app_key = 'test-key'
    client.app_secret = 'test-secret'
    client.token = 'test-token'
    calls = []
    sleeps = []
    client.sleep = sleeps.append
    def request(*args, **kwargs):
        calls.append(1)
        return {'rt_cd': '1', 'msg_cd': 'OPSQ1002'} if len(calls) < 3 else {'rt_cd': '0'}
    client._request = request
    assert client.kis('/test', 'FHKST03010230', {}) == {'rt_cd': '0'}
    assert len(calls) == 3 and sleeps == [1, 2]


def test_stale_schedule_skips_before_market_or_ledger(monkeypatch, tmp_path):
    import valuation.run as run
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None): return cls(2026, 10, 30, 15, 50, tzinfo=tz)
    monkeypatch.setattr(run, 'datetime', FrozenDateTime)
    monkeypatch.setenv('GITHUB_EVENT_NAME', 'schedule')
    monkeypatch.setenv('SCHEDULE_CRON', '47 15 * * 1-5')
    monkeypatch.setenv('TRADE_DATE', '2026-10-29')
    monkeypatch.setenv('DRY_RUN', 'false')
    monkeypatch.setattr(run, 'MarketData', lambda: pytest.fail('no market access'))
    monkeypatch.setattr(run, 'Ledger', lambda: pytest.fail('no ledger access'))
    assert run.main(['--output-dir', str(tmp_path)]) == 0
    assert json.loads((tmp_path/'run.json').read_text())['status'] == 'stale_schedule_skipped'


def test_live_delivery_requires_explicit_historical_date_and_close_time(monkeypatch):
    import valuation.run as run
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None): return cls(2026, 10, 30, 15, 0, tzinfo=tz)
    monkeypatch.setattr(run, 'datetime', FrozenDateTime)
    with pytest.raises(ValueError, match="historical"):
        run.verify_send_day(date(2026, 10, 29))
    run.verify_send_day(date(2026, 10, 29), allow_historical=True)
    with pytest.raises(ValueError, match='not been reached'):
        run.verify_send_day(date(2026, 10, 30))


def test_empty_krx_close_uses_calendar_once_to_skip_holiday(monkeypatch, tmp_path):
    import valuation.run as run
    calendar_calls = []
    class FakeMarket:
        def collect_verified_close(self, rows, day):
            raise KRXCloseEmpty('KRX confirmed close dataset empty or not yet published')
        def is_open(self, day):
            calendar_calls.append(day)
            return False
    monkeypatch.setattr(run, 'MarketData', FakeMarket)
    monkeypatch.setenv('GITHUB_EVENT_NAME', 'workflow_dispatch')
    monkeypatch.setenv('DRY_RUN', 'true')
    assert run.main(['--trade-date', '2026-10-25', '--slot', 'close',
                     '--output-dir', str(tmp_path)]) == 0
    assert calendar_calls == [date(2026, 10, 25)]
    assert json.loads((tmp_path / 'run.json').read_text())['status'] == 'market_closed'


def test_empty_krx_on_open_day_fails_without_second_price_collection(monkeypatch, tmp_path):
    import valuation.run as run
    price_calls = []
    class FakeMarket:
        def collect_verified_close(self, rows, day):
            price_calls.append(day)
            raise KRXCloseEmpty('KRX confirmed close dataset empty or not yet published')
        def is_open(self, day): return True
    monkeypatch.setattr(run, 'MarketData', FakeMarket)
    monkeypatch.setenv('GITHUB_EVENT_NAME', 'workflow_dispatch')
    monkeypatch.setenv('DRY_RUN', 'true')
    assert run.main(['--trade-date', '2026-10-30', '--slot', 'close',
                     '--output-dir', str(tmp_path)]) == 1
    assert price_calls == [date(2026, 10, 30)]
    assert 'empty or not yet published' in json.loads((tmp_path / 'run.json').read_text())['error']
